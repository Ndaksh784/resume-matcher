"""Text extraction, field parsing, ATS quality scoring and bullet coaching."""
import os
import re
from datetime import datetime

import pdfplumber
import docx

from skills_data import find_skills, SKILL_TO_CATEGORY


# --------------------------------------------------------------------------- #
# Text extraction
# --------------------------------------------------------------------------- #
def extract_text(filepath: str) -> str:
    ext = os.path.splitext(filepath)[1].lower()

    if ext == ".pdf":
        pages = []
        with pdfplumber.open(filepath) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    pages.append(page_text)
        return "\n".join(pages)

    if ext == ".docx":
        document = docx.Document(filepath)
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:  # many resume templates are tables
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)

    if ext == ".txt":
        with open(filepath, "r", encoding="utf-8", errors="ignore") as f:
            return f.read()

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX, or TXT file.")


# --------------------------------------------------------------------------- #
# Field parsing
# --------------------------------------------------------------------------- #
EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE_CANDIDATE_RE = re.compile(r"(?<![\w@])\+?\(?\d[\d\s().-]{7,}\d(?![\w@])")


def find_phone(text: str) -> str:
    """First number-like string with 10-13 digits that isn't a year range."""
    for m in PHONE_CANDIDATE_RE.finditer(text):
        digits = re.sub(r"\D", "", m.group())
        if 10 <= len(digits) <= 13 and not YEAR_RANGE_RE.search(m.group()):
            return m.group().strip()
    return ""
LINKEDIN_RE = re.compile(r"(?:https?://)?(?:www\.)?linkedin\.com/(?:in|pub)/[\w\-%]+", re.I)
GITHUB_RE = re.compile(r"(?:https?://)?(?:www\.)?github\.com/[\w\-]+", re.I)
PORTFOLIO_RE = re.compile(r"https?://(?!(?:www\.)?(?:linkedin|github)\.com)[\w.-]+\.[a-z]{2,}[^\s]*", re.I)

SECTION_ALIASES = {
    "summary": ["summary", "profile", "objective", "about me", "professional summary", "career objective"],
    "experience": ["experience", "work experience", "professional experience", "employment",
                   "internship", "internships", "work history", "employment history"],
    "education": ["education", "academic background", "academics", "qualifications", "education and training"],
    "projects": ["projects", "personal projects", "academic projects", "key projects"],
    "skills": ["skills", "technical skills", "core competencies", "key skills", "technologies", "tech stack"],
    "certifications": ["certifications", "certificates", "licenses", "courses", "licenses and certifications"],
    "achievements": ["achievements", "awards", "honors", "honours", "accomplishments", "awards and honors"],
}
HEADER_LOOKUP = {alias: key for key, aliases in SECTION_ALIASES.items() for alias in aliases}
SECTION_HEADERS = list(SECTION_ALIASES)


def split_sections(raw_text: str) -> dict:
    """Split resume text into {section_key: text} using standalone header lines."""
    buckets, current = {"header": []}, "header"
    for line in raw_text.splitlines():
        cleaned = re.sub(r"[^a-z& ]", "", line.lower()).replace("&", "and").strip()
        if cleaned in HEADER_LOOKUP and len(line.split()) <= 4:
            current = HEADER_LOOKUP[cleaned]
            buckets.setdefault(current, [])
        else:
            buckets[current].append(line)
    joined = {k: "\n".join(v).strip() for k, v in buckets.items()}
    return {k: v for k, v in joined.items() if v}


YEAR_RANGE_RE = re.compile(
    r"\b((?:19|20)\d{2})\s*(?:-|–|—|to)\s*((?:19|20)\d{2}|present|current|now|till date|ongoing)\b", re.I)


def estimate_years(sections: dict, raw_text: str) -> float:
    """Union of year ranges in the experience section (falls back to explicit claims)."""
    now = datetime.now().year
    text = sections.get("experience", "")
    spans = []
    for m in YEAR_RANGE_RE.finditer(text):
        start = int(m.group(1))
        end = now if m.group(2).isalpha() or " " in m.group(2) else int(m.group(2))
        if 1980 <= start <= end <= now + 1:
            spans.append((start, max(end, start + 0.5)))
    if spans:
        spans.sort()
        merged = [list(spans[0])]
        for s, e in spans[1:]:
            if s <= merged[-1][1]:
                merged[-1][1] = max(merged[-1][1], e)
            else:
                merged.append([s, e])
        return round(sum(e - s for s, e in merged), 1)
    claims = [int(x) for x in re.findall(r"(\d{1,2})\+?\s*(?:years|yrs)\s+(?:of\s+)?(?:\w+\s+){0,2}experience", raw_text, re.I)]
    return float(max(claims)) if claims else 0.0


def detect_education_level(text: str) -> str:
    t = text.lower()
    if re.search(r"\bph\.?d\b|doctorate", t):
        return "Doctorate"
    if re.search(r"\bm\.?\s?tech\b|\bm\.?sc\b|\bm\.?s\b|\bmba\b|\bmca\b|master", t):
        return "Master's"
    if re.search(r"\bb\.?\s?tech\b|\bb\.?e\b|\bb\.?sc\b|\bbca\b|\bb\.?s\b|bachelor", t):
        return "Bachelor's"
    if re.search(r"diploma", t):
        return "Diploma"
    return "Not detected"


def extract_fields(raw_text: str, skills_db: list = None) -> dict:
    lines = [l.strip() for l in raw_text.splitlines() if l.strip()]

    name = ""
    for line in lines[:6]:
        if EMAIL_RE.search(line) or find_phone(line) or re.search(r"https?://|www\.", line, re.I):
            continue
        if 1 <= len(line.split()) <= 5 and not re.search(r"\d", line):
            name = line.title() if line.isupper() else line
            break

    email = EMAIL_RE.search(raw_text)
    phone = find_phone(raw_text)
    linkedin = LINKEDIN_RE.search(raw_text)
    github = GITHUB_RE.search(raw_text)
    portfolio = PORTFOLIO_RE.search(raw_text)

    sections_full = split_sections(raw_text)
    sections = {k: re.sub(r"\s+", " ", sections_full[k]).strip()[:300]
                for k in SECTION_HEADERS if k in sections_full}

    skills = find_skills(raw_text)
    return {
        "name": name or "Not detected",
        "email": email.group() if email else "Not detected",
        "phone": phone or "Not detected",
        "linkedin": linkedin.group() if linkedin else "",
        "github": github.group() if github else "",
        "portfolio": portfolio.group() if portfolio else "",
        "skills": skills,
        "skills_by_category": _group(skills),
        "sections": sections,
        "years_experience": estimate_years(sections_full, raw_text),
        "education_level": detect_education_level(sections_full.get("education", raw_text)),
        "word_count": len(raw_text.split()),
    }


def _group(skills):
    grouped = {}
    for s in skills:
        grouped.setdefault(SKILL_TO_CATEGORY.get(s, "Other"), []).append(s)
    return grouped


# --------------------------------------------------------------------------- #
# Bullet coaching
# --------------------------------------------------------------------------- #
STRONG_VERBS = {
    "built", "led", "designed", "developed", "implemented", "created", "launched", "optimized",
    "reduced", "increased", "improved", "automated", "architected", "delivered", "managed", "drove",
    "engineered", "deployed", "migrated", "analyzed", "streamlined", "mentored", "scaled", "spearheaded",
    "established", "integrated", "refactored", "boosted", "generated", "negotiated", "trained", "won",
    "owned", "shipped", "simplified", "secured", "resolved", "researched", "published", "founded",
}
WEAK_OPENERS = [
    ("responsible for", ["Led", "Managed", "Owned"]),
    ("worked on", ["Built", "Developed", "Delivered"]),
    ("worked with", ["Collaborated with", "Partnered with"]),
    ("helped", ["Supported", "Enabled", "Accelerated"]),
    ("assisted", ["Supported", "Contributed to"]),
    ("duties included", ["Delivered", "Executed"]),
    ("involved in", ["Drove", "Contributed to"]),
    ("tasked with", ["Executed", "Delivered"]),
    ("participated in", ["Contributed to", "Drove"]),
]
CLICHES = ["team player", "hard working", "hardworking", "results-oriented", "results oriented",
           "detail oriented", "detail-oriented", "go-getter", "self-motivated", "synergy", "think outside the box"]


def _bullets(sections: dict, raw_text: str) -> list:
    body = "\n".join(sections.get(k, "") for k in ("experience", "projects", "achievements"))
    lines = (body if body.strip() else raw_text).splitlines()
    marker = re.compile(r"^\s*[-•*▪●◦·➢>]+\s*")
    marked = [l for l in lines if marker.match(l)]
    out = []
    for line in (marked or lines):
        clean = marker.sub("", line).strip()
        if len(clean.split()) >= 5 and not EMAIL_RE.search(clean) and not YEAR_RANGE_RE.search(clean):
            out.append(clean)
    return out


def _has_metric(text: str) -> bool:
    return bool(re.search(r"\d", re.sub(r"\b(?:19|20)\d{2}\b", "", text)))


def analyze_bullets(sections: dict, raw_text: str) -> dict:
    bullets = _bullets(sections, raw_text)
    metrics = strong = 0
    coaching = []
    for b in bullets:
        words = b.split()
        first = re.sub(r"[^a-z]", "", words[0].lower())
        low = b.lower()
        weak = next(((p, alts) for p, alts in WEAK_OPENERS if low.startswith(p)), None)
        is_strong = (first in STRONG_VERBS or first.endswith("ed")) and not weak
        has_num = _has_metric(b)
        metrics += has_num
        strong += is_strong
        if weak:
            coaching.append({"original": b[:150], "issue": f"Weak opener: “{weak[0]}”",
                             "tip": f"Start with a strong verb such as {', '.join(weak[1])}, then name the result."})
        elif not is_strong:
            coaching.append({"original": b[:150], "issue": "Doesn't start with an action verb",
                             "tip": "Open with what you did: Built, Reduced, Launched, Led…"})
        elif not has_num:
            coaching.append({"original": b[:150], "issue": "No measurable result",
                             "tip": "Add a number: % improvement, users, time saved, team size or revenue."})
    return {"total": len(bullets), "with_metrics": metrics, "strong_verbs": strong, "coaching": coaching[:6]}


def evidence_check(sections: dict) -> list:
    """Skills listed in the skills section but never shown in experience/projects."""
    listed = find_skills(sections.get("skills", ""))
    proof = "\n".join(sections.get(k, "") for k in ("experience", "projects", "achievements", "summary", "certifications"))
    if not listed or not proof.strip():
        return []
    proven = set(find_skills(proof))
    return [s for s in listed if s not in proven]


# --------------------------------------------------------------------------- #
# ATS quality score
# --------------------------------------------------------------------------- #
def calculate_quality_score(raw_text: str, fields: dict) -> dict:
    sections = split_sections(raw_text)
    bullets = analyze_bullets(sections, raw_text)
    unproven = evidence_check(sections)
    recs = []

    def rec(priority, text):
        recs.append({"priority": priority, "text": text})

    # Contact (15)
    contact = 0
    if fields["email"] != "Not detected":
        contact += 6
    else:
        rec("high", "Add a professional email address — ATS and recruiters can't reach you without it.")
    if fields["phone"] != "Not detected":
        contact += 5
    else:
        rec("high", "Add a phone number near the top of the resume.")
    if fields["name"] != "Not detected":
        contact += 4
    else:
        rec("medium", "Put your full name on its own line at the very top so it can be parsed.")

    # Sections (20)
    sec = 0
    for key, pts, msg, prio in [
        ("experience", 7, "Add a clearly labelled 'Experience' section (internships and freelance work count).", "high"),
        ("education", 5, "Add an 'Education' section with degree, institute and year.", "high"),
        ("skills", 4, "Add a dedicated 'Skills' section — it is the first place ATS looks for keywords.", "high"),
        ("projects", 2, "Add a 'Projects' section to prove your skills with real work.", "medium"),
        ("summary", 2, "Add a 2–3 line summary tailored to the role you're targeting.", "low"),
    ]:
        if key in sections:
            sec += pts
        else:
            rec(prio, msg)

    # Skills (20)
    n = len(fields["skills"])
    if n >= 12:
        skills = 20
    elif n >= 8:
        skills = 16
    elif n >= 5:
        skills = 11
        rec("medium", "Only a handful of recognised skills found. Add tools, frameworks and methods you've actually used.")
    elif n >= 3:
        skills = 6
        rec("high", "Low skill density — list more technologies and tools using their standard names.")
    else:
        skills = 2
        rec("high", "Almost no recognised skills detected. Add a Skills section using standard keywords (e.g. 'PostgreSQL', not 'databases').")

    # Impact (25)
    total = bullets["total"]
    if total:
        mr = bullets["with_metrics"] / total
        vr = bullets["strong_verbs"] / total
        metric_pts = 15 if mr >= 0.4 else 10 if mr >= 0.2 else 5 if mr > 0 else 0
        verb_pts = 10 if vr >= 0.7 else 6 if vr >= 0.4 else 2
        if mr < 0.4:
            rec("high", f"Only {bullets['with_metrics']} of {total} bullets contain numbers. Quantify results (%, users, time, money).")
        if vr < 0.7:
            rec("medium", "Start more bullets with strong action verbs (Built, Reduced, Led) instead of 'Responsible for…'.")
    else:
        metric_pts = verb_pts = 0
        rec("high", "No bullet points detected. Describe your work as short bullets that start with action verbs.")
    impact = metric_pts + verb_pts

    # Length (10)
    wc = fields["word_count"]
    if 250 <= wc <= 900:
        length = 10
    elif 150 <= wc < 250 or 900 < wc <= 1200:
        length = 6
        rec("low", "Your resume is " + ("thin — expand on achievements." if wc < 250 else "long — aim for 1–2 pages."))
    else:
        length = 3
        rec("medium", "Resume length is outside the sweet spot (roughly 250–900 words).")

    # Online presence (10)
    presence = 0
    if fields["linkedin"]:
        presence += 5
    else:
        rec("low", "Add your LinkedIn URL — recruiters check it.")
    if fields["github"] or fields["portfolio"]:
        presence += 5
    else:
        rec("low", "Add a GitHub or portfolio link so recruiters can see real work.")

    # Style hints (no points, advice only)
    low = raw_text.lower()
    if len(re.findall(r"\b(?:i|my|me)\b", low)) > 3:
        rec("low", "Avoid first-person words (I, my, me). Resumes read stronger in implied first person.")
    found_cliches = [c for c in CLICHES if c in low]
    if found_cliches:
        rec("low", f"Replace clichés like “{found_cliches[0]}” with a concrete achievement that proves it.")
    if unproven:
        rec("medium", "Skills listed but never shown in experience or projects: " + ", ".join(unproven[:6]) +
            ". Back each one with a bullet or drop it.")

    score = min(contact + sec + skills + impact + length + presence, 100)
    grade = "A" if score >= 85 else "B" if score >= 70 else "C" if score >= 55 else "D"
    order = {"high": 0, "medium": 1, "low": 2}
    recs.sort(key=lambda r: order[r["priority"]])

    return {
        "score": score,
        "grade": grade,
        "breakdown": [
            {"label": "Contact details", "score": contact, "max": 15},
            {"label": "Resume sections", "score": sec, "max": 20},
            {"label": "Skill keywords", "score": skills, "max": 20},
            {"label": "Impact & wording", "score": impact, "max": 25},
            {"label": "Length", "score": length, "max": 10},
            {"label": "Online presence", "score": presence, "max": 10},
        ],
        "feedback": recs,
        "bullets": {k: bullets[k] for k in ("total", "with_metrics", "strong_verbs")},
        "bullet_coaching": bullets["coaching"],
        "unproven_skills": unproven,
    }
